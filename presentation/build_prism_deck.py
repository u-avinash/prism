"""
Generate a concise, visual technical architecture deck for Prism.

Output: presentation/Prism_Technical_Architecture.pptx
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


# ── Visual system ────────────────────────────────────────────────────────────
W, H = Inches(13.333), Inches(7.5)
BLANK = 6

INK = RGBColor(0x0B, 0x12, 0x20)
INK_2 = RGBColor(0x12, 0x1D, 0x31)
PANEL = RGBColor(0x17, 0x26, 0x3D)
PANEL_2 = RGBColor(0x1D, 0x30, 0x4A)
WHITE = RGBColor(0xF8, 0xFA, 0xFC)
MUTED = RGBColor(0xA5, 0xB4, 0xC7)
MUTED_DARK = RGBColor(0x5C, 0x6F, 0x86)
BLUE = RGBColor(0x38, 0xBD, 0xF8)
CYAN = RGBColor(0x22, 0xD3, 0xEE)
PURPLE = RGBColor(0xA7, 0x8B, 0xFA)
PINK = RGBColor(0xF4, 0x72, 0xB6)
GREEN = RGBColor(0x34, 0xD3, 0x99)
AMBER = RGBColor(0xFB, 0xBF, 0x24)
RED = RGBColor(0xFB, 0x71, 0x85)
LINE = RGBColor(0x2B, 0x3E, 0x58)
PALE_BLUE = RGBColor(0xD9, 0xF4, 0xFF)
PALE_PURPLE = RGBColor(0xEE, 0xE9, 0xFE)
PALE_GREEN = RGBColor(0xD1, 0xFA, 0xE5)
PALE_AMBER = RGBColor(0xFE, 0xF3, 0xC7)

prs = Presentation()
prs.slide_width = W
prs.slide_height = H
blank = prs.slide_layouts[BLANK]


def rgb(hex_value):
    hex_value = hex_value.replace("#", "")
    return RGBColor(*bytes.fromhex(hex_value))


def add_slide(dark=True):
    slide = prs.slides.add_slide(blank)
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = INK if dark else WHITE
    return slide


def rect(slide, x, y, w, h, fill, line=None, radius=False, line_width=Pt(1)):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, x, y, w, h
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line:
        shape.line.color.rgb = line
        shape.line.width = line_width
    else:
        shape.line.fill.background()
    if radius:
        try:
            shape.adjustments[0] = 0.08
        except Exception:
            pass
    return shape


def text(
    slide, x, y, w, h, value, size=12, color=WHITE, bold=False,
    align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font="Aptos", italic=False,
    margin=5, line_spacing=1.0,
):
    shape = slide.shapes.add_textbox(x, y, w, h)
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Pt(margin)
    tf.margin_right = Pt(margin)
    tf.margin_top = Pt(2)
    tf.margin_bottom = Pt(2)
    tf.vertical_anchor = anchor
    for i, line in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.bold = bold
        run.font.italic = italic
    return shape


def rule(slide, x, y, w, color=BLUE, h=0.035):
    return rect(slide, x, y, w, Inches(h), color)


def line(slide, x1, y1, x2, y2, color=LINE, weight=Pt(1.5), arrow_end=False):
    con = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    con.line.color.rgb = color
    con.line.width = weight
    if arrow_end:
        con.line.end_arrowhead = True
    return con


def pill(slide, x, y, w, label, fill=PANEL_2, color=CYAN):
    rect(slide, x, y, w, Inches(0.31), fill, radius=True)
    text(slide, x, y + Inches(0.012), w, Inches(0.24), label, 8.6, color, True,
         PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)


def tag(slide, x, y, label, color=CYAN):
    w = Inches(max(0.65, len(label) * 0.073 + 0.16))
    rect(slide, x, y, w, Inches(0.26), PANEL_2, radius=True)
    text(slide, x, y + Inches(0.01), w, Inches(0.21), label.upper(), 7.5, color, True,
         PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    return w


def title(slide, eyebrow, heading, subheading=None, dark=True, page=None):
    primary = WHITE if dark else INK
    secondary = MUTED if dark else MUTED_DARK
    text(slide, Inches(0.62), Inches(0.38), Inches(6.5), Inches(0.22),
         eyebrow.upper(), 9, CYAN, True, font="Aptos Display")
    text(slide, Inches(0.6), Inches(0.67), Inches(11.9), Inches(0.52),
         heading, 27, primary, True, font="Aptos Display")
    if subheading:
        text(slide, Inches(0.62), Inches(1.22), Inches(11.6), Inches(0.32),
             subheading, 11, secondary)
    rule(slide, Inches(0.62), Inches(1.57), Inches(0.75), CYAN)
    if page:
        text(slide, Inches(12.35), Inches(0.4), Inches(0.35), Inches(0.22),
             f"{page:02d}", 9, MUTED, True, PP_ALIGN.RIGHT)


def footer(slide, dark=True):
    color = MUTED_DARK if dark else MUTED
    text(slide, Inches(0.62), Inches(7.13), Inches(8.5), Inches(0.18),
         "PRISM  /  AI-ASSISTED INCIDENT RESPONSE", 7.5, color, True)
    text(slide, Inches(10.15), Inches(7.13), Inches(2.55), Inches(0.18),
         "Technical Architecture", 7.5, color, False, PP_ALIGN.RIGHT)


def card(slide, x, y, w, h, accent, heading, body, number=None):
    rect(slide, x, y, w, h, PANEL, LINE, True)
    rect(slide, x, y, Inches(0.055), h, accent)
    if number:
        text(slide, x + Inches(0.22), y + Inches(0.16), Inches(0.38), Inches(0.28),
             number, 9, accent, True)
        heading_x = x + Inches(0.62)
        heading_w = w - Inches(0.8)
    else:
        heading_x = x + Inches(0.22)
        heading_w = w - Inches(0.42)
    text(slide, heading_x, y + Inches(0.15), heading_w, Inches(0.27),
         heading, 12, WHITE, True)
    text(slide, x + Inches(0.22), y + Inches(0.52), w - Inches(0.42), h - Inches(0.64),
         body, 9.4, MUTED, False, line_spacing=1.14)


def flow_node(slide, x, y, w, h, accent, top, body, icon):
    rect(slide, x, y, w, h, PANEL, LINE, True)
    rect(slide, x + Inches(0.16), y + Inches(0.18), Inches(0.44), Inches(0.44), accent, radius=True)
    text(slide, x + Inches(0.16), y + Inches(0.215), Inches(0.44), Inches(0.3),
         icon, 13, INK, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    text(slide, x + Inches(0.72), y + Inches(0.15), w - Inches(0.88), Inches(0.24),
         top, 10.5, WHITE, True)
    text(slide, x + Inches(0.72), y + Inches(0.42), w - Inches(0.88), h - Inches(0.52),
         body, 8.5, MUTED, False, line_spacing=1.08)


# 01 — Title
s = add_slide()
rect(s, Inches(0), Inches(0), W, H, INK)
# Decorative grid and glow-style elements
for x in [0.85, 2.15, 3.45, 4.75, 6.05, 7.35, 8.65, 9.95, 11.25, 12.55]:
    line(s, Inches(x), Inches(0), Inches(x), H, rgb("102038"), Pt(0.5))
for y in [0.9, 2.2, 3.5, 4.8, 6.1]:
    line(s, Inches(0), Inches(y), W, Inches(y), rgb("102038"), Pt(0.5))
rect(s, Inches(8.45), Inches(0.7), Inches(3.7), Inches(3.7), PANEL_2, radius=True)
rect(s, Inches(9.0), Inches(1.25), Inches(2.6), Inches(2.6), INK, radius=True)
rect(s, Inches(9.55), Inches(1.8), Inches(1.5), Inches(1.5), PANEL, radius=True)
rect(s, Inches(10.1), Inches(2.35), Inches(0.4), Inches(0.4), CYAN, radius=True)
text(s, Inches(0.78), Inches(1.25), Inches(6.3), Inches(0.25),
     "TECHNICAL ARCHITECTURE  /  2026", 11, CYAN, True)
text(s, Inches(0.72), Inches(1.72), Inches(7.8), Inches(0.82),
     "PRISM", 51, WHITE, True, font="Aptos Display")
text(s, Inches(0.78), Inches(2.72), Inches(7.2), Inches(0.8),
     "Governed AI for\nproduction incident response.", 22, rgb("D6E7F8"), False,
     font="Aptos Display", line_spacing=0.92)
rule(s, Inches(0.78), Inches(3.85), Inches(1.1), CYAN, 0.045)
text(s, Inches(0.78), Inches(4.15), Inches(6.95), Inches(0.7),
     "From OpenTelemetry signals to a reviewable RCA, proposed code fix, and pull request — with a human decision at the center.",
     13, MUTED, False, line_spacing=1.2)
pill(s, Inches(0.78), Inches(5.28), Inches(0.98), "OTEL")
pill(s, Inches(1.88), Inches(5.28), Inches(1.3), "LANGGRAPH", PANEL_2, PURPLE)
pill(s, Inches(3.3), Inches(5.28), Inches(1.14), "GITHUB", PANEL_2, GREEN)
pill(s, Inches(4.56), Inches(5.28), Inches(0.88), "JIRA", PANEL_2, AMBER)
pill(s, Inches(5.56), Inches(5.28), Inches(1.4), "HUMAN GATE", PANEL_2, PINK)
text(s, Inches(0.78), Inches(6.82), Inches(4), Inches(0.2),
     "PRISM  •  TECHNICAL ARCHITECTURE", 8, MUTED, True)


# 02 — At a glance
s = add_slide()
title(s, "The opportunity", "Compress incident response — without bypassing engineering judgement.",
      "Prism turns telemetry into an auditable, human-approved remediation package.", page=2)
metrics = [
    ("01", "INGEST", "OTLP logs & CI/CD events\nbecome project-scoped incidents", CYAN),
    ("02", "REASON", "AI-assisted severity, RCA,\nfix & test recommendation", PURPLE),
    ("03", "GOVERN", "Explicit reviewer decision\nbefore delivery actions", AMBER),
    ("04", "DELIVER", "Jira, GitHub PR & team\nnotifications after approval", GREEN),
]
for i, (num, head, body, accent) in enumerate(metrics):
    x = Inches(0.62 + i * 3.1)
    rect(s, x, Inches(2.05), Inches(2.82), Inches(2.0), PANEL, LINE, True)
    text(s, x + Inches(0.22), Inches(2.25), Inches(0.45), Inches(0.3), num, 12, accent, True)
    rect(s, x + Inches(0.22), Inches(2.72), Inches(0.52), Inches(0.052), accent)
    text(s, x + Inches(0.22), Inches(2.94), Inches(2.22), Inches(0.28), head, 14, WHITE, True)
    text(s, x + Inches(0.22), Inches(3.32), Inches(2.35), Inches(0.5), body, 9.5, MUTED)
    if i < 3:
        line(s, x + Inches(2.82), Inches(3.04), x + Inches(3.03), Inches(3.04), MUTED_DARK, Pt(1.75), True)
rect(s, Inches(0.62), Inches(4.7), Inches(12.02), Inches(1.35), PANEL_2, None, True)
text(s, Inches(0.9), Inches(4.92), Inches(1.95), Inches(0.3), "THE DESIGN PRINCIPLE", 8.5, CYAN, True)
text(s, Inches(0.9), Inches(5.25), Inches(11.25), Inches(0.46),
     "AI prepares the evidence. Humans authorize the change.", 20, WHITE, True, font="Aptos Display")
text(s, Inches(0.9), Inches(5.74), Inches(10.9), Inches(0.22),
     "No downstream Jira or GitHub delivery is initiated until an accountable reviewer approves the proposed remediation.", 9.5, MUTED)
footer(s)


# 03 — workflow
s = add_slide()
title(s, "The Prism loop", "A governed path from signal to pull request.",
      "Eight purposeful stages replace fragmented hand-offs with one visible workflow.", page=3)
steps = [
    ("01", "Signal", "OTLP log / CI event", CYAN),
    ("02", "Filter", "Parse, fingerprint & dedupe", CYAN),
    ("03", "Assess", "Severity + SLA posture", PURPLE),
    ("04", "Investigate", "Code & runtime context → RCA", PURPLE),
    ("05", "Prepare", "Fix, tests, PDF & patch", PURPLE),
    ("06", "Approve", "Human reviewer decision", AMBER),
    ("07", "Deliver", "Jira + GitHub PR", GREEN),
    ("08", "Learn", "Finalize, notify & PIR", GREEN),
]
for i, (n, h, b, accent) in enumerate(steps):
    x = Inches(0.47 + i * 1.58)
    y = Inches(2.12 if i % 2 == 0 else 3.52)
    rect(s, x + Inches(0.43), y, Inches(0.48), Inches(0.48), accent, radius=True)
    text(s, x + Inches(0.43), y + Inches(0.08), Inches(0.48), Inches(0.22), n, 9, INK, True, PP_ALIGN.CENTER)
    text(s, x, y + Inches(0.62), Inches(1.35), Inches(0.23), h, 10, WHITE, True, PP_ALIGN.CENTER)
    text(s, x - Inches(0.06), y + Inches(0.93), Inches(1.47), Inches(0.42), b, 8.3, MUTED, False, PP_ALIGN.CENTER)
    if i < 7:
        x1, y1 = x + Inches(1.18), y + Inches(0.24)
        x2 = x + Inches(1.55)
        y2 = Inches(3.0 if i % 2 == 0 else 3.0)
        line(s, x1, y1, x2, y2, LINE, Pt(1.6), True)
rect(s, Inches(4.69), Inches(5.45), Inches(4.0), Inches(0.65), PANEL_2, AMBER, True)
text(s, Inches(4.84), Inches(5.59), Inches(3.7), Inches(0.28),
     "HUMAN APPROVAL IS THE CONTROL POINT", 9, AMBER, True, PP_ALIGN.CENTER)
text(s, Inches(0.64), Inches(6.38), Inches(11.8), Inches(0.35),
     "Resilience is built in: workflow leases prevent duplicate execution, selected generation nodes retry, and unrecoverable errors are routed to a persisted FAILED state.",
     9.5, MUTED, False, PP_ALIGN.CENTER)
footer(s)


# 04 — architecture
s = add_slide()
title(s, "Architecture at a glance", "A modular platform that separates signal handling from governed remediation.",
      "Python / FastAPI services, a LangGraph state machine, project-scoped integrations, and persistent evidence.", page=4)
layers = [
    ("SIGNAL SOURCES", CYAN, ["OTLP applications", "CI/CD systems", "OTel Collector"]),
    ("INGRESS", BLUE, ["FastAPI Ingestion API  :8000", "OTLP /v1/logs  •  events  •  protected project API key"]),
    ("INTELLIGENCE", PURPLE, ["Parser + severity analysis", "SimHash deduplication", "LangGraph incident workflow"]),
    ("EXPERIENCE", PINK, ["Dashboard UI  :8080", "Review • approve • observe • administer"]),
    ("DELIVERY & DATA", GREEN, ["GitHub  •  Jira  •  Slack / Teams  •  Anypoint", "SQLite / SQLAlchemy  •  encrypted project secrets  •  PDF + patch artefacts"]),
]
for i, (label, accent, entries) in enumerate(layers):
    y = Inches(1.78 + i * 0.93)
    text(s, Inches(0.62), y + Inches(0.22), Inches(1.4), Inches(0.22), label, 8, accent, True, PP_ALIGN.RIGHT)
    rect(s, Inches(2.22), y, Inches(10.23), Inches(0.65), PANEL, LINE, True)
    rect(s, Inches(2.22), y, Inches(0.06), Inches(0.65), accent)
    if len(entries) == 1:
        text(s, Inches(2.52), y + Inches(0.19), Inches(9.55), Inches(0.24), entries[0], 11, WHITE, True)
    else:
        for j, entry in enumerate(entries):
            ex = Inches(2.52 + j * (9.34 / len(entries)))
            text(s, ex, y + Inches(0.17), Inches(9.1 / len(entries)), Inches(0.26), entry, 10.2, WHITE if i != 4 else PALE_GREEN, True, PP_ALIGN.CENTER)
            if j:
                line(s, ex - Inches(0.12), y + Inches(0.12), ex - Inches(0.12), y + Inches(0.53), LINE, Pt(1))
    if i < len(layers) - 1:
        line(s, Inches(7.34), y + Inches(0.65), Inches(7.34), y + Inches(0.86), MUTED_DARK, Pt(1.6), True)
footer(s)


# 05 — capabilities
s = add_slide()
title(s, "Capability system", "One incident record becomes a complete, reviewable engineering evidence package.",
      "Capabilities are grouped by their job in the response lifecycle.", page=5)
caps = [
    ("SEE", CYAN, "Telemetry intelligence", "OTLP parsing\nPersistent telemetry log\nFingerprinting & grouping\nLLM-first severity with fallback"),
    ("THINK", PURPLE, "Contextual diagnosis", "GitHub source context\nAnypoint runtime context\nAI-assisted RCA\nSLA and escalation posture"),
    ("PREPARE", PINK, "Remediation package", "Proposed code fix\nTest recommendation\nFix reflection / quality check\nRCA PDF + unified patch"),
    ("CONTROL", AMBER, "Accountable action", "Reviewer comments\nApprove / reject / regenerate\nWorkflow history\nAudited recovery controls"),
    ("SHIP", GREEN, "Toolchain delivery", "Jira issue & DevInfo\nGitHub branch, commit & PR\nSlack / Teams notifications\nPost-Incident Review"),
]
for i, (verb, accent, h, body) in enumerate(caps):
    x = Inches(0.62 + (i % 3) * 4.08)
    y = Inches(1.86 + (i // 3) * 2.2)
    w = Inches(3.72 if i < 3 else 5.82)
    if i == 4:
        x = Inches(6.91)
    rect(s, x, y, w, Inches(1.83), PANEL, LINE, True)
    text(s, x + Inches(0.22), y + Inches(0.18), Inches(0.86), Inches(0.23), verb, 8.5, accent, True)
    text(s, x + Inches(0.22), y + Inches(0.47), w - Inches(0.45), Inches(0.28), h, 14, WHITE, True)
    rule(s, x + Inches(0.22), y + Inches(0.84), Inches(0.52), accent, 0.025)
    text(s, x + Inches(0.22), y + Inches(1.02), w - Inches(0.45), Inches(0.6), body, 9.2, MUTED, False, line_spacing=1.12)
footer(s)


# 06 — trust
s = add_slide()
title(s, "Trust by design", "Safety is an architectural feature — not a policy afterthought.",
      "Prism protects the ingress boundary, tenant boundary, action boundary, and evidence trail.", page=6)
trust = [
    ("PROJECT-BOUND INGESTION", CYAN, "Production ingestion accepts a project API key. The resolved key, not untrusted log metadata, establishes tenant ownership."),
    ("LEAST-PRIVILEGE EXPERIENCE", PURPLE, "Admin, Team Admin, and User roles work within project-scoped access and configuration boundaries."),
    ("ENCRYPTED INTEGRATIONS", GREEN, "LLM, GitHub, Jira, notification, and Anypoint credentials are encrypted at rest using Fernet."),
    ("HUMAN DELIVERY GATE", AMBER, "A reviewer must explicitly approve a recommendation before Jira issue creation or GitHub PR automation can start."),
    ("AUDITABLE OPERATIONS", PINK, "Workflow records, bounded node events, API-key lifecycle, approval history, and recovery actions provide durable evidence."),
]
for i, (h, accent, b) in enumerate(trust):
    y = Inches(1.78 + i * 0.91)
    rect(s, Inches(0.82), y, Inches(0.55), Inches(0.55), accent, radius=True)
    text(s, Inches(0.82), y + Inches(0.115), Inches(0.55), Inches(0.23), f"0{i+1}", 9, INK, True, PP_ALIGN.CENTER)
    line(s, Inches(1.63), y + Inches(0.275), Inches(2.05), y + Inches(0.275), accent, Pt(1.5), True)
    text(s, Inches(2.2), y + Inches(0.02), Inches(3.15), Inches(0.22), h, 9, accent, True)
    text(s, Inches(2.2), y + Inches(0.26), Inches(9.65), Inches(0.3), b, 9.5, WHITE)
footer(s)


# 07 — integration ecosystem
s = add_slide()
title(s, "Integration ecosystem", "Prism works with the systems engineering teams already use.",
      "Every integration is project-scoped, with resilient client handling and encrypted configuration.", page=7)
# central core
rect(s, Inches(5.02), Inches(2.68), Inches(3.25), Inches(1.35), PANEL_2, CYAN, True, Pt(1.3))
text(s, Inches(5.25), Inches(2.91), Inches(2.8), Inches(0.26), "PRISM ORCHESTRATOR", 12, WHITE, True, PP_ALIGN.CENTER)
text(s, Inches(5.25), Inches(3.28), Inches(2.8), Inches(0.35), "LangGraph workflow\n+ integration clients", 9, MUTED, False, PP_ALIGN.CENTER)
integration_nodes = [
    (Inches(0.8), Inches(1.85), "LLM PROVIDERS", "OpenAI • Anthropic\nGemini • Groq • Ollama", PURPLE),
    (Inches(0.8), Inches(4.45), "GITHUB", "Code context\nBranch • commit • PR", GREEN),
    (Inches(9.15), Inches(1.85), "JIRA CLOUD", "Issues • comments\nDevInfo sync", AMBER),
    (Inches(9.15), Inches(4.45), "TEAM CHANNELS", "Slack • Teams\nGeneric webhook", PINK),
    (Inches(5.02), Inches(5.37), "ANYPOINT", "Runtime & deployment\ncontext", CYAN),
]
for x, y, h, b, accent in integration_nodes:
    rect(s, x, y, Inches(3.25), Inches(0.92), PANEL, LINE, True)
    rect(s, x, y, Inches(0.06), Inches(0.92), accent)
    text(s, x + Inches(0.22), y + Inches(0.16), Inches(2.8), Inches(0.2), h, 10, WHITE, True)
    text(s, x + Inches(0.22), y + Inches(0.43), Inches(2.8), Inches(0.3), b, 8.5, MUTED)
# spokes
for x1, y1, x2, y2 in [
    (4.05, 2.31, 5.02, 2.97), (4.05, 4.9, 5.02, 3.73),
    (8.27, 2.97, 9.15, 2.31), (8.27, 3.73, 9.15, 4.9), (6.65, 4.03, 6.65, 5.37)
]:
    line(s, Inches(x1), Inches(y1), Inches(x2), Inches(y2), LINE, Pt(1.5), True)
footer(s)


# 08 — experience and data
s = add_slide()
title(s, "Operational experience & evidence", "A single project-scoped view for responders, reviewers, and administrators.",
      "The dashboard makes the AI process inspectable rather than opaque.", page=8)
flow_node(s, Inches(0.65), Inches(2.0), Inches(3.55), Inches(1.2), CYAN, "INCIDENT COMMAND", "Incident list, severity, status,\nSLA posture, grouping & telemetry", "◈")
flow_node(s, Inches(4.89), Inches(2.0), Inches(3.55), Inches(1.2), PURPLE, "REVIEW WORKSPACE", "RCA, code fix, test idea, patch,\nPDF report, comments & decision", "✓")
flow_node(s, Inches(9.13), Inches(2.0), Inches(3.55), Inches(1.2), GREEN, "ADMIN CONSOLE", "Projects, teams, integrations,\nAPI-key lifecycle & audit evidence", "⚙")
for x in [4.2, 8.44]:
    line(s, x, Inches(2.6), x + Inches(0.65), Inches(2.6), LINE, Pt(1.7), True)
rect(s, Inches(0.65), Inches(4.3), Inches(12.03), Inches(1.28), PANEL_2, LINE, True)
text(s, Inches(0.94), Inches(4.55), Inches(2.2), Inches(0.23), "PERSISTED EVIDENCE", 9, CYAN, True)
text(s, Inches(3.08), Inches(4.45), Inches(8.9), Inches(0.28),
     "Incident record  →  workflow history  →  approval record  →  delivery links", 14, WHITE, True, PP_ALIGN.CENTER)
text(s, Inches(1.16), Inches(5.01), Inches(10.95), Inches(0.23),
     "SQLite / SQLAlchemy stores incidents and telemetry; encrypted project integration settings are separated from operational workflow records; RCA PDFs and patch files are retained as artefacts.",
     9, MUTED, False, PP_ALIGN.CENTER)
footer(s)


# 09 — runtime reality
s = add_slide()
title(s, "Production readiness: current foundation, next evolution", "A pragmatic single-instance implementation with a clear enterprise path.",
      "The distinction matters: Prism is deployable today, while scale-out capabilities should be planned deliberately.", page=9)
columns = [
    ("CURRENT FOUNDATION", GREEN, [
        ("FastAPI services", "Dedicated ingestion (:8000) and UI (:8080) processes"),
        ("Single source of truth", "SQLite + SQLAlchemy, filesystem artefacts"),
        ("Operational controls", "API keys, sessions, RBAC, audit and workflow leases"),
        ("Provider resilience", "Retry/backoff and multi-provider LLM fallback"),
    ]),
    ("EVOLUTION PATH", AMBER, [
        ("Scale & availability", "PostgreSQL and a worker / queue model as volume grows"),
        ("Repeatable delivery", "Container image, IaC and automated deployment pipeline"),
        ("Telemetry maturity", "Persist and correlate OTLP traces & metrics"),
        ("Enterprise perimeter", "TLS gateway, distributed limiting, OIDC / SSO as required"),
    ]),
]
for i, (h, accent, rows) in enumerate(columns):
    x = Inches(0.7 + i * 6.08)
    rect(s, x, Inches(1.88), Inches(5.86), Inches(4.75), PANEL, LINE, True)
    rect(s, x, Inches(1.88), Inches(5.86), Inches(0.55), accent, radius=True)
    text(s, x, Inches(2.03), Inches(5.86), Inches(0.22), h, 10, INK, True, PP_ALIGN.CENTER)
    for j, (head, body) in enumerate(rows):
        y = Inches(2.72 + j * 0.88)
        text(s, x + Inches(0.25), y, Inches(2.1), Inches(0.2), head, 9.5, accent, True)
        text(s, x + Inches(2.33), y, Inches(3.2), Inches(0.38), body, 8.8, MUTED)
        if j < 3:
            line(s, x + Inches(0.25), y + Inches(0.58), x + Inches(5.55), y + Inches(0.58), LINE, Pt(0.7))
footer(s)


# 10 — close
s = add_slide()
rect(s, Inches(0), Inches(0), W, H, INK)
rect(s, Inches(0), Inches(0), W, Inches(0.1), CYAN)
text(s, Inches(0.8), Inches(1.0), Inches(5.5), Inches(0.22), "THE PRISM OUTCOME", 10, CYAN, True)
text(s, Inches(0.74), Inches(1.43), Inches(10.7), Inches(1.12),
     "A faster path to\nresponsible remediation.", 31, WHITE, True, font="Aptos Display", line_spacing=0.92)
text(s, Inches(0.8), Inches(3.0), Inches(7.15), Inches(0.52),
     "Prism brings observability, AI-assisted engineering analysis, accountable approval, and delivery automation together in one governed workflow.",
     13, MUTED, False, line_spacing=1.18)
outcomes = [
    ("LESS TRIAGE", "Filter and prioritize noisy telemetry before costly investigation.", CYAN),
    ("BETTER EVIDENCE", "Give reviewers an RCA, fix proposal, tests, PDF and patch — not just an alert.", PURPLE),
    ("CONTROLLED ACTION", "Preserve human accountability before any engineering-system change is initiated.", AMBER),
]
for i, (head, body, accent) in enumerate(outcomes):
    y = Inches(4.17 + i * 0.65)
    rect(s, Inches(0.82), y, Inches(0.12), Inches(0.12), accent, radius=True)
    text(s, Inches(1.17), y - Inches(0.04), Inches(1.65), Inches(0.22), head, 9.2, accent, True)
    text(s, Inches(2.88), y - Inches(0.05), Inches(6.35), Inches(0.3), body, 9.4, WHITE)
rect(s, Inches(9.75), Inches(1.18), Inches(2.4), Inches(4.85), PANEL_2, None, True)
text(s, Inches(10.03), Inches(1.6), Inches(1.84), Inches(0.3), "SIGNAL", 11, CYAN, True, PP_ALIGN.CENTER)
line(s, Inches(10.95), Inches(2.12), Inches(10.95), Inches(2.56), CYAN, Pt(2), True)
text(s, Inches(10.03), Inches(2.68), Inches(1.84), Inches(0.3), "REASON", 11, PURPLE, True, PP_ALIGN.CENTER)
line(s, Inches(10.95), Inches(3.2), Inches(10.95), Inches(3.64), PURPLE, Pt(2), True)
text(s, Inches(10.03), Inches(3.76), Inches(1.84), Inches(0.3), "APPROVE", 11, AMBER, True, PP_ALIGN.CENTER)
line(s, Inches(10.95), Inches(4.28), Inches(10.95), Inches(4.72), AMBER, Pt(2), True)
text(s, Inches(10.03), Inches(4.84), Inches(1.84), Inches(0.3), "DELIVER", 11, GREEN, True, PP_ALIGN.CENTER)
text(s, Inches(0.8), Inches(6.84), Inches(5.8), Inches(0.18),
     "PRISM  /  TECHNICAL ARCHITECTURE", 8, MUTED, True)


# Save
out_path = Path(__file__).resolve().parent / "Prism_Technical_Architecture.pptx"
prs.save(out_path)
print(f"Saved: {out_path}")
print(f"Slides: {len(prs.slides)}")
