"""Professional PPTX presentation generator for Prism."""
from __future__ import annotations

from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor as PPTColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches as PI, Pt as PPt, Emu

NAVY = PPTColor(11, 16, 32)
PANEL = PPTColor(21, 28, 44)
PURPLE = PPTColor(124, 92, 252)
CYAN = PPTColor(20, 184, 166)
AMBER = PPTColor(242, 184, 75)
WHITE = PPTColor(244, 247, 255)
MUTED = PPTColor(167, 176, 197)
FONT = "Calibri"


def _bg(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY


def _txbox(slide, text, x, y, w, h, size=11, bold=False, color=WHITE, align=PP_ALIGN.LEFT):
    tf = slide.shapes.add_textbox(PI(x), PI(y), PI(w), PI(h)).text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = PPt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = FONT


def _rule(slide, x, y, w, color=PURPLE, thickness=3000):
    from pptx.util import Pt
    line = slide.shapes.add_shape(1, PI(x), PI(y), PI(w), Emu(thickness))
    line.fill.solid()
    line.fill.fore_color.rgb = color
    line.line.fill.background()


def _bullets_box(slide, items, x, y, w, h, size=17, color=WHITE, heading=None, hcolor=None):
    tf = slide.shapes.add_textbox(PI(x), PI(y), PI(w), PI(h)).text_frame
    tf.word_wrap = True
    first = True
    if heading:
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run()
        r.text = heading
        r.font.size = PPt(size + 3)
        r.font.bold = True
        r.font.color.rgb = hcolor or CYAN
        r.font.name = FONT
        first = False
    for item in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run()
        r.text = f"  {item}"
        r.font.size = PPt(size)
        r.font.color.rgb = color
        r.font.name = FONT
        p.space_before = PPt(4)


def _img(slide, path: Path, x, y, w):
    slide.shapes.add_picture(str(path), PI(x), PI(y), width=PI(w))


def _slide_header(slide, title, subtitle=None):
    _txbox(slide, title, 0.45, 0.22, 12.5, 0.65, size=26, bold=True, color=WHITE)
    if subtitle:
        _txbox(slide, subtitle, 0.45, 0.82, 12.5, 0.4, size=12, color=MUTED)
    _rule(slide, 0.45, 1.18, 12.4, PURPLE)


def build(out_path: Path, assets: Path, record: dict, events: list[tuple]) -> Path:
    prs = Presentation()
    prs.slide_width = PI(13.333)
    prs.slide_height = PI(7.5)
    blank = prs.slide_layouts[6]

    # ------------------------------------------------------------------ Slide 1: Title
    s = prs.slides.add_slide(blank); _bg(s)
    _rule(s, 0.6, 2.8, 0.15, PURPLE, 50000)
    _txbox(s, "Prism", 1.0, 1.6, 11.3, 1.2, size=54, bold=True, color=WHITE)
    _txbox(s, "AI-Powered Incident Management Platform", 1.0, 2.85, 11.3, 0.7, size=24, color=CYAN)
    _txbox(s, "From telemetry signal to human-approved remediation", 1.0, 3.55, 11.3, 0.5, size=16, color=MUTED)
    _txbox(s, "Author: Upadhyayula Avinash  |  October 2026  |  NTT DATA", 1.0, 6.7, 11.3, 0.5, size=11, color=MUTED)

    # ------------------------------------------------------------------ Slide 2: Problem Statement
    s = prs.slides.add_slide(blank); _bg(s)
    _slide_header(s, "The Operational Challenge", "Why existing workflows break under production pressure")
    problems = [
        "Production errors demand immediate, structured investigation -- but teams are overwhelmed.",
        "Logs, source context, ticketing, notifications, and Pull Requests are fragmented, manual steps.",
        "On-call engineers repeat the same investigative patterns for every incident.",
        "There is no consistent, auditable chain from a raw error to a safe, reviewed remediation.",
        "The result: slower MTTR, inconsistent quality, and accountability gaps.",
    ]
    _bullets_box(s, problems, 0.6, 1.45, 12.1, 5.5, size=18)

    # ------------------------------------------------------------------ Slide 3: What is Prism
    s = prs.slides.add_slide(blank); _bg(s)
    _slide_header(s, "What is Prism?", "The platform that closes the gap between alert and remedy")
    what = [
        "Receives OpenTelemetry log signals from MuleSoft/Anypoint or any compatible application.",
        "Classifies severity, detects technology context, and deduplicates repeated errors automatically.",
        "Uses a configurable LLM to produce a Root Cause Analysis (RCA) with a confidence score.",
        "Generates a concrete code fix, a quality reflection, an RCA PDF report, and a unified patch.",
        "Pauses for a human reviewer to approve, reject, or request a revised fix.",
        "After approval: sends notifications, creates a Jira ticket, opens a GitHub Pull Request.",
    ]
    _bullets_box(s, what, 0.6, 1.45, 8.5, 5.5, size=17)
    _txbox(s, "Automation accelerates\nevidence gathering.\nPeople decide what ships.", 9.5, 2.0, 3.4, 3.2, size=18, bold=True, color=CYAN, align=PP_ALIGN.CENTER)

    # ------------------------------------------------------------------ Slide 4: Architecture
    s = prs.slides.add_slide(blank); _bg(s)
    _slide_header(s, "Architecture Overview", "Two services, one AI workflow, one shared data store")
    _img(s, assets / "01_prism_architecture.png", 0.45, 1.35, 12.45)

    # ------------------------------------------------------------------ Slide 5: 11-Step Workflow
    s = prs.slides.add_slide(blank); _bg(s)
    _slide_header(s, "The 11-Step LangGraph Workflow", "Retryable, durable, and pause-safe by design")
    steps_pre = [
        "1. assess_severity  -- Classify from OTLP signal; fire 'Incident Created' notification",
        "2. generate_rca  -- LLM RCA with confidence score; fetches live GitHub/Anypoint context",
        "3. generate_fix  -- LLM-generated code fix with explanation",
        "4. generate_pdf  -- RCA PDF report saved to data/pdfs/",
        "5. reflect  -- Quality scores: correctness, safety, code quality, completeness",
        "6. generate_patch  -- Unified .patch file saved to data/patches/",
        "7. await_approval  -- PAUSE: human reviews RCA, fix diff, and patch before deciding",
    ]
    steps_post = [
        "8. send_notifications  -- Slack / Teams webhooks (from per-project DB config)",
        "9. create_jira  -- Jira ticket, patch attachment, branch/commit/PR synced to DevInfo",
        "10. create_pr  -- GitHub branch, commit fixed file, open Pull Request",
        "11. finalize  -- Persist final status, timestamps, and all integration artifact URLs",
    ]
    _bullets_box(s, steps_pre, 0.55, 1.38, 6.5, 5.2, size=13.5, color=WHITE, heading="Pre-approval (automated)", hcolor=CYAN)
    _bullets_box(s, steps_post, 6.9, 1.38, 6.0, 5.2, size=13.5, color=WHITE, heading="Post-approval (automated)", hcolor=AMBER)

    # ------------------------------------------------------------------ Slide 6: Incident Journey
    s = prs.slides.add_slide(blank); _bg(s)
    _slide_header(s, "The Incident Journey in Six Stages", "Prism accelerates the journey; humans govern the outcome")
    _img(s, assets / "02_incident_journey.png", 0.45, 1.35, 12.45)

    # ------------------------------------------------------------------ Slide 7: Capabilities
    s = prs.slides.add_slide(blank); _bg(s)
    _slide_header(s, "Core Capabilities", "Everything a team needs from alert to remedy")
    col1 = [
        "OTLP intake with technology detection",
        "SimHash error deduplication",
        "Severity classification (HIGH / CRITICAL)",
        "LLM-generated RCA with confidence score",
        "LLM-generated code fix and explanation",
        "Quality reflection (4-axis scoring)",
    ]
    col2 = [
        "Auto-generated RCA PDF report",
        "Downloadable unified diff patch",
        "Approve / reject / regenerate with feedback",
        "Post-approval: Slack/Teams notification",
        "Post-approval: Jira ticket + patch attach",
        "Post-approval: GitHub branch + commit + PR",
    ]
    col3 = [
        "Durable workflow runs and step events",
        "Governed post-approval recovery",
        "Multi-tenant project isolation",
        "Per-project encrypted credentials",
        "Role-based access (Admin / TA / User)",
        "Security audit trail and CSV export",
    ]
    _bullets_box(s, col1, 0.5, 1.38, 4.1, 5.5, size=15)
    _bullets_box(s, col2, 4.7, 1.38, 4.1, 5.5, size=15)
    _bullets_box(s, col3, 9.0, 1.38, 4.1, 5.5, size=15)

    # ------------------------------------------------------------------ Slide 8: Role model
    s = prs.slides.add_slide(blank); _bg(s)
    _slide_header(s, "Role, Tenant & Project Model", "Clean separation of platform, configuration, and review concerns")
    _img(s, assets / "04_role_model.png", 0.45, 1.35, 12.45)

    # ------------------------------------------------------------------ Slide 9: Security
    s = prs.slides.add_slide(blank); _bg(s)
    _slide_header(s, "Security, Governance & Reliability", "Fail-closed by default; every action is auditable")
    _img(s, assets / "06_security_governance.png", 0.45, 1.35, 12.45)

    # ------------------------------------------------------------------ Slide 10: Technology
    s = prs.slides.add_slide(blank); _bg(s)
    _slide_header(s, "Technology Stack", "Pragmatic, proven, and provider-swappable")
    _img(s, assets / "05_technology_stack.png", 0.45, 1.35, 12.45)

    # ------------------------------------------------------------------ Slide 11: Verified example
    s = prs.slides.add_slide(blank); _bg(s)
    _slide_header(s, "Verified Local Example: Incident NTTD6XI2", "Real evidence from the local Prism database")
    facts = [
        f"Application: {record['app_name']}  |  Environment: {record['environment']}",
        f"Severity: {record['severity']}  |  Status: {record['status']}  |  Technology: MuleSoft / DataWeave",
        f"RCA confidence: {float(record['rca_confidence']) * 100:.1f}%  |  Approval: {record.get('approval_status', 'n/a')}",
        f"Jira ticket: {record.get('jira_ticket_key', 'n/a')}  |  GitHub PR: #{record.get('pr_number', 'n/a')}",
        "",
        "Root cause: DataWeave coercion of a missing orderDate field to Date type caused CRITICAL",
        "production failure.  Fix: null-safe existence check with explicit date format.",
        "",
        "Full run: RCA generated in 165 s  |  Fix generated in 178 s  |  Patch in 0.06 s",
        "Post-approval workflow completed in 65 s  |  Total: ~7.5 minutes end-to-end.",
    ]
    _bullets_box(s, facts, 0.55, 1.38, 8.0, 5.5, size=15)
    _img(s, assets / "03_verified_incident_timeline.png", 8.5, 1.38, 4.45)

    # ------------------------------------------------------------------ Slide 12: Integration model
    s = prs.slides.add_slide(blank); _bg(s)
    _slide_header(s, "Integration Model", "Every integration is per-project, encrypted, and optional")
    integrations = [
        ("LLM Provider", "OpenAI, Anthropic, Azure OpenAI, Gemini, Groq, or local Ollama", CYAN),
        ("GitHub", "Code fetch, branch/commit/PR creation", PURPLE),
        ("Jira", "Ticket creation, patch attach, DevInfo sync", CYAN),
        ("Slack / Teams", "Incident and workflow event notifications", PURPLE),
        ("MuleSoft Anypoint", "Live CloudHub/ARM runtime context for richer RCA", CYAN),
    ]
    y = 1.55
    for name, desc, color in integrations:
        _txbox(s, name, 0.6, y, 2.8, 0.38, size=15, bold=True, color=color)
        _txbox(s, desc, 3.5, y, 9.5, 0.38, size=14, color=WHITE)
        _rule(s, 0.6, y + 0.42, 12.2, PANEL, 4000)
        y += 0.9

    # ------------------------------------------------------------------ Slide 13: Product boundary
    s = prs.slides.add_slide(blank); _bg(s)
    _slide_header(s, "Current Boundary & Roadmap", "Honest about what Prism does and does not do today")
    current = [
        "OTLP logs and derived incidents are fully persisted and rendered.",
        "Workflow runs, step events, and security audit records are durable and queryable.",
        "RCA, fix, quality reflection, PDF, and patch are generated when LLM is configured.",
        "Post-approval Jira/GitHub/Slack require valid per-project integration credentials.",
    ]
    roadmap = [
        "Trace and metric signals are acknowledged today; full persistence is a planned next step.",
        "Fix proposals require human review; auto-merge is not on the near-term roadmap.",
        "Expanded CI/CD feedback loops and static-analysis signals to improve RCA confidence.",
    ]
    _bullets_box(s, current, 0.55, 1.45, 6.0, 3.0, size=16, heading="What Prism does today", hcolor=CYAN)
    _bullets_box(s, roadmap, 0.55, 4.2, 6.0, 2.8, size=16, heading="Near-term planned improvements", hcolor=AMBER)
    _txbox(s, "Prism is designed to be honest\nabout its boundaries -- no silent\nfailures, no fabricated results.", 7.5, 2.2, 5.5, 3.0, size=18, bold=True, color=CYAN, align=PP_ALIGN.CENTER)

    # ------------------------------------------------------------------ Slide 14: Adoption path
    s = prs.slides.add_slide(blank); _bg(s)
    _slide_header(s, "Adoption Path", "From zero to a governed, production-ready deployment")
    steps_adopt = [
        "1. Deploy Prism with production security settings and a durable, shared encryption key.",
        "2. As admin: create a project and assign a Team Admin.",
        "3. As Team Admin: configure LLM, then optionally GitHub, Jira, Slack/Teams.",
        "4. Issue a project ingestion key; configure collector/gateway to forward OTLP to Prism.",
        "5. Run a controlled first incident: verify RCA quality, approval flow, and downstream delivery.",
        "6. Measure review time, workflow success rate, and remediation quality; tune thresholds.",
        "7. Expand to additional projects and teams with confidence.",
    ]
    _bullets_box(s, steps_adopt, 0.55, 1.45, 12.1, 5.5, size=17)

    # ------------------------------------------------------------------ Slide 15: Close
    s = prs.slides.add_slide(blank); _bg(s)
    _rule(s, 0.6, 2.6, 0.15, CYAN, 50000)
    _txbox(s, "Prism", 1.0, 1.5, 11.3, 1.0, size=44, bold=True, color=WHITE)
    _txbox(s, "Autonomous evidence gathering. Human accountability.", 1.0, 2.65, 11.3, 0.65, size=22, color=CYAN)
    _txbox(s, "Automate what machines do well. Reserve decisions for people.", 1.0, 3.35, 11.3, 0.5, size=16, color=MUTED)
    _txbox(s, "GitHub: https://github.com/u-avinash  |  LinkedIn: https://www.linkedin.com/in/avinash-upadhyayula/", 1.0, 6.6, 11.3, 0.5, size=11, color=MUTED)

    prs.save(str(out_path))
    return out_path
