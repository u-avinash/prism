"""High-Level Design Word document generator for Prism."""
from __future__ import annotations

from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT

ACCENT = RGBColor(124, 92, 252)
ACCENT2 = RGBColor(20, 184, 166)
DARK = RGBColor(11, 16, 32)
GREY = RGBColor(90, 98, 120)
FONT = "Calibri"


def _title_page(doc, title, subtitle, meta_lines):
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(40)
    r.font.color.rgb = ACCENT
    r.font.name = FONT

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(subtitle)
    r.font.size = Pt(18)
    r.font.color.rgb = DARK
    r.font.name = FONT

    doc.add_paragraph()
    for line in meta_lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line)
        r.font.size = Pt(11)
        r.font.color.rgb = GREY
        r.font.name = FONT
    doc.add_page_break()


def _heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for r in p.runs:
        r.font.name = FONT
        r.font.color.rgb = ACCENT if level == 1 else DARK
    return p


def _para(doc, text, size=11, color=None, bold=False, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.name = FONT
    r.bold = bold
    r.italic = italic
    if color:
        r.font.color.rgb = color
    return p


def _bullets(doc, items, style="List Bullet"):
    for item in items:
        p = doc.add_paragraph(item, style=style)
        for r in p.runs:
            r.font.name = FONT
            r.font.size = Pt(11)


def _table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, text in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(10.5)
        r.font.name = FONT
        r.font.color.rgb = RGBColor(255, 255, 255)
        hdr[i]._tc.get_or_add_tcPr()
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            r = p.add_run(str(text))
            r.font.size = Pt(10)
            r.font.name = FONT
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph()
    return table


def _figure(doc, path: Path, caption: str, width=6.6):
    doc.add_picture(str(path), width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(caption)
    r.italic = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = GREY
    r.font.name = FONT


def build(doc_path: Path, assets: Path, record: dict, events: list[tuple]) -> Path:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    _title_page(
        doc,
        "Prism",
        "High-Level Design Document",
        [
            "AI-Powered, Human-Governed Incident Management Platform",
            "Version 1.0  |  October 2026",
            "Author: Upadhyayula Avinash  |  Prepared for engineering, operations, and leadership audiences",
            "Classification: Internal / Partner-shareable (contains no live credentials or customer data)",
        ],
    )

    # ---------------- Table of contents (manual) ----------------
    _heading(doc, "Table of Contents", level=1)
    toc_items = [
        "1. Executive Summary",
        "2. Business Problem and Motivation",
        "3. Goals and Non-Goals",
        "4. System Overview and Architecture",
        "5. End-to-End Incident Journey",
        "6. Core Capabilities",
        "7. Component Deep Dive",
        "8. Role, Tenant, and Project Model",
        "9. Technology Stack",
        "10. Security, Governance, and Reliability",
        "11. Data Model and Persistence",
        "12. Integration Model",
        "13. Current Product Boundaries and Roadmap",
        "14. Verified Local Example: Incident NTTD6XI2",
        "15. Deployment and Operations Guide",
        "16. Glossary",
    ]
    _bullets(doc, toc_items, style="List Number")
    doc.add_page_break()

    # ---------------- 1. Executive Summary ----------------
    _heading(doc, "1. Executive Summary")
    _para(
        doc,
        "Prism is an AI-powered incident management platform built for engineering and operations "
        "teams who need production errors investigated quickly, consistently, and safely. Instead of "
        "manually piecing together logs, source code, tickets, and notifications after every failure, "
        "Prism automates the repetitive investigative work and uses a Large Language Model (LLM) to "
        "propose a Root Cause Analysis (RCA) and a candidate code fix. A human reviewer always makes "
        "the final call before anything is delivered downstream.",
    )
    _para(
        doc,
        "In plain terms: when an application throws a serious error, Prism explains what probably went "
        "wrong, suggests a specific code change to fix it, and prepares a reviewable package -- a PDF "
        "report and a patch file -- so an engineer can approve, reject, or request changes with one "
        "click. Once approved, Prism can automatically notify the team, open a Jira ticket, and create a "
        "GitHub Pull Request with the fix.",
    )
    _figure(doc, assets / "01_prism_architecture.png", "Figure 1-1. Prism high-level architecture.")

    # ---------------- 2. Business problem ----------------
    _heading(doc, "2. Business Problem and Motivation")
    _para(
        doc,
        "Modern production systems generate large volumes of telemetry, but the distance between "
        "'an error happened' and 'a safe fix is ready for review' is usually long and manual. Typical "
        "friction includes:",
    )
    _bullets(doc, [
        "Engineers must manually correlate error logs with the responsible service and code location.",
        "Root-cause investigation is repetitive and time-consuming, especially under on-call pressure.",
        "Creating a ticket, writing a fix, and opening a Pull Request are separate, disconnected steps.",
        "Teams often lack a consistent, auditable trail connecting a production error to its remediation.",
    ])
    _para(
        doc,
        "Prism closes this gap by acting as an always-available first responder: it gathers evidence, "
        "proposes a diagnosis and fix, and prepares the downstream paperwork -- while keeping a human "
        "in control of what actually ships.",
    )

    # ---------------- 3. Goals and non-goals ----------------
    _heading(doc, "3. Goals and Non-Goals")
    _heading(doc, "3.1 Goals", level=2)
    _bullets(doc, [
        "Reduce the time between error occurrence and a reviewable root-cause explanation.",
        "Generate a concrete, explainable code-level fix proposal, not just a generic suggestion.",
        "Preserve human accountability: no code is merged, and no ticket created, without approval.",
        "Operate safely across multiple clients/projects (multi-tenant) with isolated credentials.",
        "Provide durable operational evidence (workflow runs, step events, audit log) for governance.",
    ])
    _heading(doc, "3.2 Non-Goals", level=2)
    _bullets(doc, [
        "Prism does not auto-merge or auto-deploy code changes without human approval.",
        "Prism does not replace a full APM/observability platform for traces and metrics today.",
        "Prism does not guarantee a correct fix; it is an assistant whose output must be reviewed.",
    ])

    # ---------------- 4. System overview ----------------
    _heading(doc, "4. System Overview and Architecture")
    _para(
        doc,
        "Prism is composed of two cooperating FastAPI services, a stateful AI workflow engine, and a "
        "SQLite-backed persistence layer. The Ingestion API focuses purely on receiving telemetry; the "
        "Dashboard UI focuses purely on human interaction, review, and administration. This separation "
        "keeps the machine-facing surface small and auditable, while the human-facing surface can evolve "
        "independently.",
    )
    _figure(doc, assets / "01_prism_architecture.png", "Figure 4-1. Component architecture and data flow.")
    _table(
        doc,
        ["Layer", "Technology", "Default Port", "Responsibility"],
        [
            ("Ingestion API", "FastAPI + OpenTelemetry", "8000", "Receive/validate OTLP telemetry; create incidents"),
            ("Dashboard UI", "FastAPI + Jinja2 + SSE", "8080", "Review, approve, administer, observe"),
            ("Workflow Engine", "LangGraph", "n/a", "Stateful multi-step AI incident processing"),
            ("Database", "SQLite via SQLAlchemy", "n/a", "Durable incident, telemetry, and audit records"),
            ("LLM Layer", "Pluggable providers", "n/a", "RCA generation, fix generation, quality reflection"),
        ],
        col_widths=[1.4, 2.0, 0.9, 2.3],
    )

    # ---------------- 5. End-to-end journey ----------------
    _heading(doc, "5. End-to-End Incident Journey")
    _para(
        doc,
        "Every incident that meets the configured severity threshold (HIGH or CRITICAL by default) "
        "moves through the same six-stage journey. Steps 1 through 4 run automatically after ingestion. "
        "Step 5 is a deliberate pause point. Step 6 only executes after a human approves the work.",
    )
    _figure(doc, assets / "02_incident_journey.png", "Figure 5-1. The six-stage incident journey.")
    _table(
        doc,
        ["Stage", "What happens", "Who/what acts"],
        [
            ("1. Receive", "An OTLP log record arrives at the Ingestion API.", "Source application / OTel Collector"),
            ("2. Understand", "Parse, detect technology, deduplicate, classify severity.", "Prism Ingestion API"),
            ("3. Analyze", "Generate RCA with a confidence score, using code/runtime context.", "LangGraph + LLM"),
            ("4. Recommend", "Propose a fix, run a quality reflection, build PDF + patch.", "LangGraph + LLM"),
            ("5. Review", "Pause for human approval, rejection, or regeneration request.", "Reviewer (User role)"),
            ("6. Deliver", "Notify Slack/Teams, create Jira ticket, open GitHub PR.", "Prism post-approval workflow"),
        ],
        col_widths=[1.1, 3.6, 2.0],
    )
    _para(
        doc,
        "This boundary -- automation up to a recommendation, a human decision in the middle, and "
        "automation again after approval -- is the central design choice of Prism. It is what allows "
        "the system to move quickly without removing accountability.",
    )

    # ---------------- 6. Core capabilities ----------------
    _heading(doc, "6. Core Capabilities")
    _bullets(doc, [
        "OTLP log ingestion with technology detection (MuleSoft, Java, Python, Node.js, .NET, Go, Ruby, PHP, Rust, Kotlin/JVM, and more).",
        "SimHash-based error deduplication so repeated identical failures do not create duplicate incidents.",
        "Configurable severity classification (HIGH / CRITICAL triggers full automated analysis).",
        "LLM-generated Root Cause Analysis with a self-assessed confidence score.",
        "LLM-generated code fix proposal with a human-readable explanation of the change.",
        "Automated quality reflection scoring correctness, safety, code quality, and completeness.",
        "Automatically generated RCA PDF report and a downloadable unified diff patch file.",
        "One-click human approval, rejection, or feedback-driven fix regeneration.",
        "Post-approval automation: Slack/Teams notification, Jira ticket creation with patch attachment, GitHub branch + commit + Pull Request.",
        "Durable workflow-run and step-event history for diagnosis, auditing, and governed recovery.",
        "Multi-tenant project isolation with per-project encrypted integration credentials.",
        "Role-based access control across Admin, Team Admin, and User personas.",
    ])

    # ---------------- 7. Component deep dive ----------------
    _heading(doc, "7. Component Deep Dive")
    _heading(doc, "7.1 Ingestion API (ingestion/)", level=2)
    _para(doc, "Receives OTLP/HTTP log payloads, parses them into structured incident and telemetry records, "
                "applies deduplication and severity analysis, persists results, and triggers the LangGraph "
                "workflow as a background task so the HTTP response is never blocked on AI processing.")
    _heading(doc, "7.2 Workflow Engine (agents/)", level=2)
    _para(doc, "A LangGraph state machine coordinates severity assessment, RCA generation, fix generation, "
                "PDF/patch creation, quality reflection, the approval pause, and -- after approval -- "
                "notification and Jira/GitHub delivery. Retryable nodes automatically retry transient "
                "failures before routing to a terminal FAILED state.")
    _heading(doc, "7.3 Dashboard UI (ui/)", level=2)
    _para(doc, "A FastAPI + Jinja2 application serving incident lists and detail pages, live workflow progress "
                "via Server-Sent Events, approval controls, observability views, and the Admin / Team Admin "
                "consoles for project and integration management.")
    _heading(doc, "7.4 Integration Clients (integrations/)", level=2)
    _para(doc, "Thin, dedicated clients for the LLM provider abstraction, GitHub, Jira, Slack/Teams webhooks, "
                "and MuleSoft Anypoint, each configured per project from encrypted stored credentials.")
    _heading(doc, "7.5 Persistence Layer (storage/)", level=2)
    _para(doc, "SQLAlchemy models and repositories manage incidents, telemetry logs, workflow runs and step "
                "events, security audit events, and encrypted integration configuration, backed by SQLite.")

    # ---------------- 8. Role model ----------------
    _heading(doc, "8. Role, Tenant, and Project Model")
    _para(doc, "Prism is multi-tenant: every incident, telemetry record, and integration credential belongs to "
                "exactly one Project. This keeps clients, teams, or business units cleanly separated.")
    _figure(doc, assets / "04_role_model.png", "Figure 8-1. Role hierarchy and responsibilities.")
    _table(
        doc,
        ["Role", "Scope", "Representative capabilities"],
        [
            ("admin", "Platform-wide", "Create/delete projects and users; assign Team Admins"),
            ("team_admin", "Project-scoped", "Configure integrations; manage team members; issue ingestion keys"),
            ("user", "Project-scoped", "View incidents; approve/reject fixes; add review comments"),
        ],
        col_widths=[1.2, 1.3, 4.2],
    )

    # ---------------- 9. Technology stack ----------------
    _heading(doc, "9. Technology Stack")
    _figure(doc, assets / "05_technology_stack.png", "Figure 9-1. Technology stack by functional area.")
    _table(
        doc,
        ["Area", "Technology Choices"],
        [
            ("Language & services", "Python 3.11+, FastAPI, Uvicorn"),
            ("Agent orchestration", "LangGraph, LangChain"),
            ("Web experience", "Jinja2 templates, JavaScript, Server-Sent Events, CSS"),
            ("Data persistence", "SQLite, SQLAlchemy"),
            ("AI providers", "OpenAI, Anthropic, Azure OpenAI, Google Gemini, Groq, or local Ollama"),
            ("Telemetry", "OpenTelemetry OTLP/HTTP; optional OpenTelemetry Collector"),
            ("Delivery integrations", "GitHub REST API, Jira REST/DevInfo API, Slack/Teams webhooks, MuleSoft Anypoint"),
            ("Artifacts & resilience", "ReportLab/PyPDF for PDFs, Tenacity for retries, Fernet for credential encryption"),
            ("Quality & CI", "pytest, GitHub Actions compile/test gate"),
        ],
        col_widths=[1.7, 5.0],
    )

    # ---------------- 10. Security & governance ----------------
    _heading(doc, "10. Security, Governance, and Reliability")
    _figure(doc, assets / "06_security_governance.png", "Figure 10-1. Security and governance pillars.")
    _bullets(doc, [
        "Tenant isolation: incidents, telemetry, and duplicate detection are scoped per project.",
        "Credential protection: all integration credentials are encrypted at rest with a Fernet key.",
        "Ingestion key lifecycle: plaintext keys are shown once at issuance; use and revocation are audited.",
        "Fail-closed production defaults: authenticated ingestion, assigned-project requirement, secure "
        "cookies, CSRF protection, and explicit CORS origins are enforced; insecure overrides are rejected.",
        "Durable workflow evidence: every run and node-level event is recorded for diagnosis and recovery.",
        "Governed recovery: only admin/team_admin can trigger post-approval recovery, with a mandatory "
        "operator reason, a persistent execution lease, and an audit trail entry.",
    ])

    # ---------------- 11. Data model ----------------
    _heading(doc, "11. Data Model and Persistence")
    _para(doc, "The SQLite database at data/incidents.db holds all durable operational state:")
    _table(
        doc,
        ["Table", "Purpose"],
        [
            ("incidents", "Core incident record: workflow state, RCA, fix, approval, integration artefacts"),
            ("telemetry_logs", "Raw OTLP log records, independent of incident creation"),
            ("project_integration_configs", "Encrypted per-project credentials and ingestion-key digests"),
            ("security_audit_events", "Append-only lifecycle events for keys and privileged actions"),
            ("workflow_runs", "Execution lifecycle for full, post-approval, and recovery workflows"),
            ("workflow_step_events", "Bounded operational evidence for each workflow node attempt"),
            ("incident_comments", "Reviewer comment threads on incidents"),
        ],
        col_widths=[2.1, 4.6],
    )

    # ---------------- 12. Integration model ----------------
    _heading(doc, "12. Integration Model")
    _para(doc, "Every integration is configured per project by a Team Admin and stored encrypted. Prism "
                "degrades gracefully when an integration is not configured -- for example, RCA generation "
                "still runs without GitHub access, simply with less code context.")
    _table(
        doc,
        ["Integration", "Purpose"],
        [
            ("LLM provider", "RCA, fix, and reflection generation (OpenAI, Anthropic, Azure OpenAI, Gemini, Groq, Ollama)"),
            ("GitHub", "Source code retrieval; branch, commit, and Pull Request creation"),
            ("Jira", "Ticket creation, patch attachment, and Development panel sync"),
            ("Slack / Microsoft Teams", "Incident and workflow event notifications"),
            ("MuleSoft Anypoint", "Live CloudHub/ARM runtime context for richer RCA"),
        ],
        col_widths=[1.8, 4.9],
    )

    # ---------------- 13. Boundaries & roadmap ----------------
    _heading(doc, "13. Current Product Boundaries and Roadmap")
    _para(doc, "Transparency about current limitations is part of Prism's design philosophy:")
    _bullets(doc, [
        "OTLP trace and metric endpoints validate and acknowledge payloads today, but raw traces and "
        "metrics are not yet persisted or rendered in the dashboard.",
        "GitHub, Jira, Slack, Teams, Anypoint, and LLM outcomes all require valid, correctly scoped "
        "per-project configuration; without it, Prism will not fabricate results.",
        "Fix proposals are recommendations. They are reviewed for correctness, safety, and domain fit "
        "before being merged -- Prism does not claim certainty.",
    ])
    _para(doc, "Natural next steps include persisted trace/metric correlation, expanded static-analysis "
                "signals feeding severity classification, and deeper CI/CD feedback loops for fix validation.")

    doc.add_page_break()

    # ---------------- 14. Verified local example ----------------
    _heading(doc, "14. Verified Local Example: Incident NTTD6XI2")
    _para(doc, "To ground this document in real behaviour rather than a hypothetical, the local Prism "
                "database was queried for a genuine, fully processed incident. The following details are "
                "read directly from local Prism records; no production system was accessed to produce them.")
    _table(
        doc,
        ["Field", "Recorded value"],
        [
            ("Incident ID", record["incident_id"]),
            ("Application", record["app_name"]),
            ("Environment", record["environment"]),
            ("Severity", record["severity"]),
            ("Final status", record["status"]),
            ("Technology", "MuleSoft / DataWeave"),
            ("RCA confidence", f"{float(record['rca_confidence']) * 100:.1f}%" if record.get("rca_confidence") else "n/a"),
            ("Approval status", record.get("approval_status") or "n/a"),
            ("Repository", record.get("repo_full_name") or "n/a"),
            ("Error location", f"{record.get('error_file_path')}:{record.get('error_line_number')}"),
            ("Jira ticket", record.get("jira_ticket_key") or "n/a"),
            ("GitHub Pull Request", f"#{record.get('pr_number')}" if record.get("pr_number") else "n/a"),
        ],
        col_widths=[2.0, 4.7],
    )
    _figure(doc, assets / "03_verified_incident_timeline.png", "Figure 14-1. Actual recorded workflow timeline for NTTD6XI2.")
    _para(doc, "In this incident, a DataWeave transform attempted to coerce a missing orderDate field to a "
                "Date, causing a CRITICAL production failure in order-processing-service. Prism generated an "
                "RCA with 87.5% confidence, proposed a null-safe fix, produced a PDF report and a unified "
                "patch, and recorded human approval before the post-approval workflow completed successfully, "
                "referencing Jira SCRUM-12 and GitHub Pull Request #61.")
    _para(doc, "The complete evidence package for this incident -- including the actual generated patch file "
                "and RCA PDF -- is available at docs/incident-evidence/NTTD6XI2/ alongside this document.",
          italic=True)

    # ---------------- 15. Deployment guide ----------------
    _heading(doc, "15. Deployment and Operations Guide")
    _heading(doc, "15.1 Local Quick Start", level=2)
    _bullets(doc, [
        "Install dependencies: pip install -r requirements.txt",
        "Start the Ingestion API: uvicorn ingestion.api:app --host 0.0.0.0 --port 8000",
        "Start the Dashboard UI: uvicorn ui.server:app --host 0.0.0.0 --port 8080",
        "Log in at http://localhost:8080 with the seeded first-boot admin credentials.",
        "As admin, onboard a project and assign a Team Admin; as Team Admin, configure integrations.",
    ])
    _heading(doc, "15.2 Production Hardening Checklist", level=2)
    _bullets(doc, [
        "Set PRISM_ENVIRONMENT=production and review every fail-closed setting before go-live.",
        "Terminate TLS at a reverse proxy; expose only HTTPS publicly.",
        "Use a durable, externally managed INTEGRATION_SECRET_KEY shared by all processes.",
        "Route OTLP through a managed collector or API gateway with a distinct project API key.",
        "Change the seeded administrator password immediately after first login.",
    ])
    _heading(doc, "15.3 Operational Monitoring", level=2)
    _bullets(doc, [
        "GET /api/workflow-health -- project-scoped success rate, active runs, average duration.",
        "GET /api/incidents/{id}/workflow-runs -- execution history for a specific incident.",
        "GET /api/workflow-runs/{run_id}/events -- chronological node-level operational evidence.",
        "/audit page -- security audit trail for ingestion-key and privileged operator actions.",
    ])

    # ---------------- 16. Glossary ----------------
    _heading(doc, "16. Glossary")
    _table(
        doc,
        ["Term", "Meaning"],
        [
            ("OTLP", "OpenTelemetry Protocol -- the standard format Prism uses to receive logs/traces/metrics"),
            ("RCA", "Root Cause Analysis -- the AI-generated explanation of why an error occurred"),
            ("LangGraph", "A framework for building stateful, multi-step AI agent workflows"),
            ("Patch file", "A unified diff describing the exact proposed code change"),
            ("Workflow run", "One durable, auditable execution of the incident-processing state machine"),
            ("Severity threshold", "The minimum severity (HIGH or CRITICAL) that triggers full AI processing"),
            ("Tenant / Project", "An isolated unit of ownership for incidents, data, and credentials"),
        ],
        col_widths=[1.6, 5.1],
    )
    _para(doc, "End of document. Generated from the verified Prism source tree and local incident NTTD6XI2 "
                "on 2026-10-14. This document contains no live credentials and does not claim access to "
                "external GitHub, Jira, or Slack systems.", italic=True, color=GREY)

    doc.save(str(doc_path))
    return doc_path
